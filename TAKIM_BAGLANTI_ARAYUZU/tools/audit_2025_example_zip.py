from __future__ import annotations

import csv
import json
import shutil
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from typing import Any


PROJECT_ROOT = Path(__file__).resolve().parents[1]

ZIP_PATH = Path(
    r"C:\Users\muham\OneDrive\Masaüstü\HYZ_2025_Ornek_Veriler.zip"
)

OUTPUT_DIR = PROJECT_ROOT / "outputs" / "example_data_audit"
EXTRACTED_SMALL_FILES_DIR = OUTPUT_DIR / "extracted_small_files"

VIDEO_EXTENSIONS = {
    ".mp4",
    ".avi",
    ".mov",
    ".mkv",
    ".mpeg",
    ".mpg",
}

TABLE_EXTENSIONS = {
    ".csv",
    ".xlsx",
    ".xls",
    ".tsv",
}

TEXT_EXTENSIONS = {
    ".txt",
    ".json",
    ".yaml",
    ".yml",
    ".xml",
}

SAFE_EXTRACT_EXTENSIONS = TABLE_EXTENSIONS | TEXT_EXTENSIONS

MAX_SAFE_FILE_SIZE_MB = 100
MAX_SAFE_FILE_SIZE_BYTES = MAX_SAFE_FILE_SIZE_MB * 1024 * 1024
MAX_PREVIEW_ROWS = 10


def human_size(size_bytes: int) -> str:
    value = float(size_bytes)

    for unit in ("B", "KB", "MB", "GB", "TB"):
        if value < 1024.0 or unit == "TB":
            return f"{value:.2f} {unit}"
        value /= 1024.0

    return f"{size_bytes} B"


def normalize_zip_path(name: str) -> Path:
    clean_name = name.replace("\\", "/").lstrip("/")
    path = Path(clean_name)

    safe_parts = [
        part
        for part in path.parts
        if part not in ("", ".", "..")
    ]

    return Path(*safe_parts)


def safe_extract_member(
    archive: zipfile.ZipFile,
    member: zipfile.ZipInfo,
    destination_root: Path,
) -> Path:
    relative_path = normalize_zip_path(member.filename)
    destination = destination_root / relative_path

    resolved_root = destination_root.resolve()
    resolved_destination = destination.resolve()

    if resolved_root not in resolved_destination.parents:
        raise ValueError(
            f"Güvensiz ZIP yolu engellendi: {member.filename}"
        )

    destination.parent.mkdir(parents=True, exist_ok=True)

    with archive.open(member, "r") as source:
        with destination.open("wb") as target:
            shutil.copyfileobj(source, target)

    return destination


def read_text_preview(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "type": "text",
        "encoding": None,
        "line_count_previewed": 0,
        "preview": [],
        "error": None,
    }

    encodings = (
        "utf-8-sig",
        "utf-8",
        "cp1254",
        "latin-1",
    )

    for encoding in encodings:
        try:
            with path.open(
                "r",
                encoding=encoding,
                errors="strict",
            ) as file:
                lines = []

                for _, line in zip(range(30), file):
                    lines.append(line.rstrip("\r\n"))

            result["encoding"] = encoding
            result["line_count_previewed"] = len(lines)
            result["preview"] = lines
            return result

        except UnicodeDecodeError:
            continue
        except Exception as exc:
            result["error"] = f"{type(exc).__name__}: {exc}"
            return result

    result["error"] = "Dosya uygun encoding ile okunamadı."
    return result


def detect_csv_delimiter(sample: str) -> str:
    candidates = [",", ";", "\t", "|"]

    try:
        dialect = csv.Sniffer().sniff(
            sample,
            delimiters="".join(candidates),
        )
        return dialect.delimiter
    except csv.Error:
        return ","


def read_csv_preview(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "type": "csv",
        "encoding": None,
        "delimiter": None,
        "columns": [],
        "rows": [],
        "error": None,
    }

    encodings = (
        "utf-8-sig",
        "utf-8",
        "cp1254",
        "latin-1",
    )

    for encoding in encodings:
        try:
            text = path.read_text(encoding=encoding)
            sample = text[:8192]
            delimiter = detect_csv_delimiter(sample)

            reader = csv.reader(
                text.splitlines(),
                delimiter=delimiter,
            )

            rows = []

            for index, row in enumerate(reader):
                rows.append(row)

                if index >= MAX_PREVIEW_ROWS:
                    break

            result["encoding"] = encoding
            result["delimiter"] = (
                "\\t" if delimiter == "\t" else delimiter
            )

            if rows:
                result["columns"] = rows[0]
                result["rows"] = rows[1:]

            return result

        except UnicodeDecodeError:
            continue
        except Exception as exc:
            result["error"] = f"{type(exc).__name__}: {exc}"
            return result

    result["error"] = "CSV uygun encoding ile okunamadı."
    return result


def read_excel_preview(path: Path) -> dict[str, Any]:
    result: dict[str, Any] = {
        "type": "excel",
        "sheets": {},
        "error": None,
    }

    try:
        from openpyxl import load_workbook
    except ImportError:
        result["error"] = (
            "openpyxl kurulu değil. Excel içeriği okunmadı; "
            "dosya yine de çıkarıldı."
        )
        return result

    if path.suffix.lower() == ".xls":
        result["error"] = (
            "Eski .xls formatı openpyxl ile okunamaz."
        )
        return result

    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

        for sheet_name in workbook.sheetnames:
            worksheet = workbook[sheet_name]
            rows = []

            for row_index, row in enumerate(
                worksheet.iter_rows(values_only=True)
            ):
                rows.append(
                    [
                        None if value is None else str(value)
                        for value in row
                    ]
                )

                if row_index >= MAX_PREVIEW_ROWS:
                    break

            result["sheets"][sheet_name] = {
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "preview_rows": rows,
            }

        workbook.close()

    except Exception as exc:
        result["error"] = f"{type(exc).__name__}: {exc}"

    return result


def preview_file(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()

    if suffix in {".csv", ".tsv"}:
        return read_csv_preview(path)

    if suffix in {".xlsx", ".xls"}:
        return read_excel_preview(path)

    return read_text_preview(path)


def build_text_report(report: dict[str, Any]) -> str:
    lines: list[str] = []

    lines.append("=" * 90)
    lines.append("HYZ 2025 ÖRNEK VERİ ZIP AUDIT")
    lines.append("=" * 90)
    lines.append(f"Tarih: {report['audit_time']}")
    lines.append(f"ZIP: {report['zip_path']}")
    lines.append(f"ZIP boyutu: {report['zip_size_human']}")
    lines.append(
        f"Toplam ZIP girdisi: {report['total_entries']}"
    )
    lines.append(
        f"Toplam sıkıştırılmamış boyut: "
        f"{report['total_uncompressed_human']}"
    )
    lines.append("")

    lines.append("[VİDEOLAR]")

    if report["videos"]:
        for video in report["videos"]:
            lines.append(
                f"- {video['name']} | "
                f"sıkıştırılmış={video['compressed_human']} | "
                f"gerçek={video['uncompressed_human']}"
            )
    else:
        lines.append("Video bulunamadı.")

    lines.append("")
    lines.append("[ÇIKARILAN KÜÇÜK DOSYALAR]")

    if report["extracted_files"]:
        for file_info in report["extracted_files"]:
            lines.append(
                f"- {file_info['archive_name']} -> "
                f"{file_info['output_path']} | "
                f"{file_info['size_human']}"
            )
    else:
        lines.append("Çıkarılan dosya bulunamadı.")

    lines.append("")
    lines.append("[ATLANAN DOSYALAR]")

    if report["skipped_files"]:
        for file_info in report["skipped_files"]:
            lines.append(
                f"- {file_info['name']} | "
                f"{file_info['reason']} | "
                f"{file_info['size_human']}"
            )
    else:
        lines.append("Atlanan dosya yok.")

    lines.append("")
    lines.append("[GÜVENLİK]")

    lines.append("- Videolar çıkarılmadı.")
    lines.append("- ZIP değiştirilmedi.")
    lines.append("- main.py çalıştırılmadı.")
    lines.append("- Sunucu bağlantısı kurulmadı.")
    lines.append("- Model dosyaları değiştirilmedi.")
    lines.append("- .env okunmadı.")

    return "\n".join(lines)


def main() -> None:
    print("=" * 80)
    print("HYZ 2025 ÖRNEK VERİ ZIP AUDIT")
    print("=" * 80)
    print(f"ZIP yolu: {ZIP_PATH}")
    print()

    if not ZIP_PATH.is_file():
        print("[HATA] ZIP dosyası bulunamadı.")
        print("Dosya yolunu kontrol et:")
        print(ZIP_PATH)
        sys.exit(1)

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    EXTRACTED_SMALL_FILES_DIR.mkdir(
        parents=True,
        exist_ok=True,
    )

    report: dict[str, Any] = {
        "audit_time": datetime.now().isoformat(
            timespec="seconds"
        ),
        "zip_path": str(ZIP_PATH),
        "zip_size_bytes": ZIP_PATH.stat().st_size,
        "zip_size_human": human_size(
            ZIP_PATH.stat().st_size
        ),
        "total_entries": 0,
        "total_uncompressed_bytes": 0,
        "total_uncompressed_human": None,
        "videos": [],
        "extracted_files": [],
        "skipped_files": [],
        "previews": {},
        "errors": [],
    }

    with zipfile.ZipFile(ZIP_PATH, "r") as archive:
        members = [
            member
            for member in archive.infolist()
            if not member.is_dir()
        ]

        report["total_entries"] = len(members)
        report["total_uncompressed_bytes"] = sum(
            member.file_size for member in members
        )
        report["total_uncompressed_human"] = human_size(
            report["total_uncompressed_bytes"]
        )

        for member in members:
            suffix = Path(member.filename).suffix.lower()

            entry = {
                "name": member.filename,
                "compressed_bytes": member.compress_size,
                "compressed_human": human_size(
                    member.compress_size
                ),
                "uncompressed_bytes": member.file_size,
                "uncompressed_human": human_size(
                    member.file_size
                ),
            }

            if suffix in VIDEO_EXTENSIONS:
                report["videos"].append(entry)
                print(
                    f"[VİDEO] {member.filename} | "
                    f"{entry['uncompressed_human']}"
                )
                continue

            if suffix not in SAFE_EXTRACT_EXTENSIONS:
                report["skipped_files"].append(
                    {
                        "name": member.filename,
                        "reason": (
                            f"Desteklenmeyen uzantı: "
                            f"{suffix or 'uzantısız'}"
                        ),
                        "size_bytes": member.file_size,
                        "size_human": human_size(
                            member.file_size
                        ),
                    }
                )
                continue

            if member.file_size > MAX_SAFE_FILE_SIZE_BYTES:
                report["skipped_files"].append(
                    {
                        "name": member.filename,
                        "reason": (
                            f"{MAX_SAFE_FILE_SIZE_MB} MB sınırını "
                            "aşıyor"
                        ),
                        "size_bytes": member.file_size,
                        "size_human": human_size(
                            member.file_size
                        ),
                    }
                )
                continue

            try:
                extracted_path = safe_extract_member(
                    archive=archive,
                    member=member,
                    destination_root=(
                        EXTRACTED_SMALL_FILES_DIR
                    ),
                )

                extracted_info = {
                    "archive_name": member.filename,
                    "output_path": str(extracted_path),
                    "size_bytes": member.file_size,
                    "size_human": human_size(
                        member.file_size
                    ),
                }

                report["extracted_files"].append(
                    extracted_info
                )

                report["previews"][member.filename] = (
                    preview_file(extracted_path)
                )

                print(
                    f"[ÇIKARILDI] {member.filename} | "
                    f"{human_size(member.file_size)}"
                )

            except Exception as exc:
                error_message = (
                    f"{member.filename}: "
                    f"{type(exc).__name__}: {exc}"
                )

                report["errors"].append(error_message)
                print(f"[HATA] {error_message}")

    json_report_path = OUTPUT_DIR / "zip_audit_report.json"
    text_report_path = OUTPUT_DIR / "zip_audit_report.txt"

    json_report_path.write_text(
        json.dumps(
            report,
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )

    text_report_path.write_text(
        build_text_report(report),
        encoding="utf-8",
    )

    print()
    print("=" * 80)
    print("AUDIT TAMAMLANDI")
    print("=" * 80)
    print(f"Video sayısı          : {len(report['videos'])}")
    print(
        f"Çıkarılan küçük dosya : "
        f"{len(report['extracted_files'])}"
    )
    print(
        f"Atlanan dosya         : "
        f"{len(report['skipped_files'])}"
    )
    print(f"Hata                   : {len(report['errors'])}")
    print()
    print(f"Metin raporu: {text_report_path}")
    print(f"JSON raporu : {json_report_path}")
    print()
    print("Videolar çıkarılmadı.")
    print("Ana sistem değiştirilmedi.")
    print("main.py çalıştırılmadı.")
    print("Sunucu bağlantısı kurulmadı.")


if __name__ == "__main__":
    main()